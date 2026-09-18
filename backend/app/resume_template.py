"""Template-preserving resume editing.

If the user uploaded a .docx, it IS the template: we detect its sections and
styles, let the LLM rewrite content per section, then rebuild the file cloning
the original formatting — same order, same look, new words.

PDFs can't be faithfully re-laid-out, so for those we return the rewrite in
the user's own section order plus an explicit change list.
"""
from __future__ import annotations

import re
from copy import deepcopy
from io import BytesIO


def is_heading_like(text: str, style_name: str = "") -> bool:
    s = (text or "").strip()
    if not s or len(s) > 60:
        return False
    if style_name.lower().startswith("heading"):
        return True
    if s.isupper() and len(s.split()) <= 5:
        return True
    return False


def analyze_docx(blob: bytes) -> dict:
    from docx import Document

    doc = Document(BytesIO(blob))
    sections: list[dict] = []
    cur = {"heading": "HEADER", "paras": []}
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        if is_heading_like(t, p.style.name if p.style else "") and cur["paras"]:
            sections.append(cur)
            cur = {"heading": t, "paras": []}
        elif is_heading_like(t, p.style.name if p.style else "") and not cur["paras"] and cur["heading"] != "HEADER":
            sections.append(cur)
            cur = {"heading": t, "paras": []}
        elif is_heading_like(t, p.style.name if p.style else "") and cur["heading"] == "HEADER" and not cur["paras"]:
            sections.append(cur)
            cur = {"heading": t, "paras": []}
        else:
            cur["paras"].append(t)
    if cur["paras"] or cur["heading"] != "HEADER":
        sections.append(cur)
    return {"type": "docx", "sections": [s["heading"] for s in sections if s["heading"] != "HEADER"],
            "count": len(sections)}


def _copy_run_fmt(src, dst) -> None:
    try:
        dst.bold = src.bold
        dst.italic = src.italic
        if src.font.size:
            dst.font.size = src.font.size
        if src.font.name:
            dst.font.name = src.font.name
        if src.font.color and src.font.color.rgb:
            dst.font.color.rgb = src.font.color.rgb
    except Exception:
        pass


def _set_para_text(para, text: str) -> None:
    style = para.style
    fmt_src = None
    for r in para.runs:
        if r.text.strip():
            fmt_src = r
            break
    if fmt_src is None and para.runs:
        fmt_src = para.runs[0]
    p_el, runs = para._p, list(para.runs)
    for r in runs:
        p_el.remove(r._r)
    run = para.add_run(text)
    if fmt_src is not None:
        _copy_run_fmt(fmt_src, run)
    try:
        para.style = style
    except Exception:
        pass


def rebuild_docx(template_blob: bytes, new_sections: dict[str, list[str]]) -> bytes:
    """Rewrite content section-by-section, cloning original styles."""
    from docx import Document
    from docx.oxml import OxmlElement

    norm = {re.sub(r"\W+", "", k).lower(): v for k, v in (new_sections or {}).items()}
    doc = Document(BytesIO(template_blob))
    paras = doc.paragraphs
    # index headings
    heads: list[tuple[int, str]] = []
    for i, p in enumerate(paras):
        t = p.text.strip()
        if t and is_heading_like(t, p.style.name if p.style else ""):
            heads.append((i, t))
    for hi, (pi, title) in enumerate(heads):
        key = re.sub(r"\W+", "", title).lower()
        if key not in norm:
            continue
        new_paras = [t for t in norm[key] if t.strip()]
        start = pi + 1
        end = heads[hi + 1][0] if hi + 1 < len(heads) else len(paras)
        body_idx = [j for j in range(start, end) if paras[j].text.strip()]
        # replace in place
        for k, j in enumerate(body_idx):
            if k < len(new_paras):
                _set_para_text(paras[j], new_paras[k])
            else:
                _set_para_text(paras[j], "")
        # append extras cloning last body style
        if len(new_paras) > len(body_idx) and body_idx:
            anchor = paras[body_idx[-1]]._p
            for extra in new_paras[len(body_idx):]:
                src = paras[body_idx[-1]]
                new_p = OxmlElement("w:p")
                anchor.addnext(new_p)
                from docx.text.paragraph import Paragraph as _P

                p_obj = _P(new_p, doc)
                try:
                    p_obj.style = src.style
                except Exception:
                    pass
                run = p_obj.add_run(extra)
                if src.runs:
                    _copy_run_fmt(src.runs[0], run)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def save_upload(data_dir: str, user_id: str, filename: str, blob: bytes) -> str:
    import os as _os

    safe = "".join(c for c in user_id if c.isalnum() or c in "@._-")[:60] or "anon"
    d = _os.path.join(data_dir, "resumes", safe)
    _os.makedirs(d, exist_ok=True)
    ext = (filename.rsplit(".", 1)[-1].lower() if "." in (filename or "") else "bin")[:5]
    if ext not in ("pdf", "docx", "txt"):
        ext = "bin"
    p = _os.path.join(d, f"original.{ext}")
    with open(p, "wb") as f:
        f.write(blob)
    return p


def load_upload(data_dir: str, user_id: str) -> tuple[str, bytes] | tuple[None, None]:
    import glob as _glob
    import os as _os

    safe = "".join(c for c in user_id if c.isalnum() or c in "@._-")[:60] or "anon"
    hits = _glob.glob(_os.path.join(data_dir, "resumes", safe, "original.*"))
    if not hits:
        return None, None
    with open(hits[0], "rb") as f:
        return hits[0], f.read()
